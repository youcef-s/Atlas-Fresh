from collections.abc import Callable
from pathlib import Path

import openpyxl
import pytest

from app.main import DEFAULT_WORKBOOK
from app.models import Snapshot
from app.validation import load_snapshot

Mutator = Callable[[openpyxl.Workbook], None]


@pytest.fixture
def baseline() -> Snapshot:
    return load_snapshot(DEFAULT_WORKBOOK)


@pytest.fixture
def mutated(tmp_path: Path) -> Callable[[Mutator], Path]:
    """Copy the source workbook, apply an edit, and return the new path (source stays untouched)."""

    def make(edit: Mutator) -> Path:
        wb = openpyxl.load_workbook(DEFAULT_WORKBOOK)
        edit(wb)
        out = tmp_path / "mutated.xlsx"
        wb.save(out)
        return out

    return make


def cell(wb: openpyxl.Workbook, sheet: str, row_id: str, column: str) -> openpyxl.cell.Cell:
    """Locate a data cell by its ID in column A and its header name."""
    ws = wb[sheet]
    header_row = next(
        r for r in ws.iter_rows() if r[0].value in ("farm_id", "client_id", "station_id")
    )
    col = next(c.column for c in header_row if c.value == column)
    row = next(r for r in ws.iter_rows() if r[0].value == row_id)
    return ws.cell(row=row[0].row, column=col)
